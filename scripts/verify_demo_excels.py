import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

def verify_student_file(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    print(f"\n--- Checking {path} ---")
    print(f"Total rows: {len(rows)}")
    header = rows[4]
    print(f"Header row (row 5): {header}")
    data_rows = rows[5:]
    print(f"Data rows count: {len(data_rows)}")
    for i, r in enumerate(data_rows):
        stt, code, name, dob, gender, class_code, email, phone, addr = r[:9]
        if i in [4, 9, 14, 19]:
            print(f"  Row {i+1} [STT {stt}]: code={code}, class={class_code}, email={email}")

def verify_grade_file(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    print(f"\n--- Checking {path} ---")
    print(f"Total rows: {len(rows)}")
    header = rows[4]
    print(f"Header row (row 5): {header}")
    data_rows = rows[5:]
    print(f"Data rows count: {len(data_rows)}")
    for i, r in enumerate(data_rows):
        enr_id, code, name, cc1, cc2, mid, fin = r[:7]
        if i in [4, 9, 14, 19]:
            print(f"  Row {i+1}: enr_id={enr_id}, code={code}, cc1={cc1}, cc2={cc2}, mid={mid}, fin={fin}")

verify_student_file("demo_excel/01_SinhVien_Moi_HopLe.xlsx")
verify_student_file("demo_excel/02_SinhVien_CoLoi_KiemTraDoiSoat.xlsx")
verify_grade_file("demo_excel/03_BangDiem_HopLe_IT301.xlsx")
verify_grade_file("demo_excel/04_BangDiem_CoLoi_Validation.xlsx")
