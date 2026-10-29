import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = os.path.join(os.getcwd(), "demo_excel")
os.makedirs(OUTPUT_DIR, exist_ok=True)

FONT_NAME = "Calibri"

title_font = Font(name=FONT_NAME, size=15, bold=True, color="1E3A8A")
subtitle_font = Font(name=FONT_NAME, size=11, bold=True, color="1F2937")
note_font = Font(name=FONT_NAME, size=10, italic=True, color="4B5563")

header_fill_blue = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Dark Blue
header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

highlight_error_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light red
highlight_error_font = Font(name=FONT_NAME, size=11, bold=True, color="B91C1C")

thin_border_side = Side(style="thin", color="CBD5E1")
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

# Danh sách 25 sinh viên mẫu phòng thi P.01 (Môn IT301 - Lập trình Web)
SAMPLE_STUDENTS = [
    ("IT301-P01-001", "24010201", "Nguyễn Hải Đăng", "15/03/2004", "K19-SE1", 8.5, None, ""),
    ("IT301-P01-002", "24010202", "Trần Thu Hà", "22/07/2004", "K19-SE1", 9.0, None, ""),
    ("IT301-P01-003", "24010203", "Lê Hoàng Nam", "10/11/2004", "K19-SE1", 7.5, None, ""),
    ("IT301-P01-004", "24010204", "Phạm Minh Tuấn", "05/01/2004", "K19-SE1", 8.0, None, ""),
    ("IT301-P01-005", "24010205", "Vũ Phương Linh", "18/09/2004", "K19-SE1", 9.5, None, ""),
    ("IT301-P01-006", "24010206", "Hoàng Gia Bảo", "30/04/2004", "K19-SE1", 6.5, None, ""),
    ("IT301-P01-007", "24010207", "Đặng Ngọc Ánh", "12/12/2004", "K19-SE1", 8.0, None, ""),
    ("IT301-P01-008", "24010208", "Bùi Quốc Anh", "25/08/2004", "K19-SE1", 7.0, None, ""),
    ("IT301-P01-009", "24010209", "Đỗ Thanh Hằng", "19/02/2004", "K19-SE1", 8.5, None, ""),
    ("IT301-P01-010", "24010210", "Ngô Đình Trọng", "08/06/2004", "K19-SE1", 0.0, "V", "Vắng thi không phép"),
    ("IT301-P01-011", "24010211", "Dương Thùy Trang", "14/10/2004", "K19-SE1", 9.0, None, ""),
    ("IT301-P01-012", "24010212", "Lý Quang Diệu", "27/05/2004", "K19-SE1", 8.0, None, ""),
    ("IT301-P01-013", "24010213", "Mai Bảo Ngọc", "03/03/2004", "K19-SE1", 8.5, None, ""),
    ("IT301-P01-014", "24010214", "Đinh Hữu Phước", "16/08/2004", "K19-SE1", 7.5, None, ""),
    ("IT301-P01-015", "24010215", "Lâm Mỹ Uyên", "29/11/2004", "K19-SE1", None, "I", "Hoãn thi do ốm có đơn"),
    ("IT301-P01-016", "24010216", "Trịnh Văn Quyết", "01/07/2004", "K19-SE1", 7.0, None, ""),
    ("IT301-P01-017", "24010217", "Hà Kiều Anh", "11/09/2004", "K19-SE1", 8.5, None, ""),
    ("IT301-P01-018", "24010218", "Lưu Đức Hải", "20/04/2004", "K19-SE1", 9.0, None, ""),
    ("IT301-P01-019", "24010219", "Tạ Bích Loan", "07/12/2004", "K19-SE1", 8.0, None, ""),
    ("IT301-P01-020", "24010220", "Phan Thanh Hùng", "15/06/2004", "K19-SE1", 7.5, None, ""),
    ("IT301-P01-021", "24010221", "Châu Vĩnh Khang", "23/01/2004", "K19-SE1", 8.5, None, ""),
    ("IT301-P01-022", "24010222", "Quách Tuấn Du", "04/10/2004", "K19-SE1", 8.0, None, ""),
    ("IT301-P01-023", "24010223", "Ân Bích Trâm", "17/05/2004", "K19-SE1", 9.0, None, ""),
    ("IT301-P01-024", "24010224", "Vương Đình Huệ", "30/08/2004", "K19-SE1", 8.5, None, ""),
    ("IT301-P01-025", "24010225", "Thái Thùy Linh", "09/02/2004", "K19-SE1", 9.5, None, ""),
]

# =============================================================
# FILE 5: 05_PhongThi_HopLe_IT301_P01.xlsx (Happy Path)
# =============================================================
wb5 = openpyxl.Workbook()
ws5 = wb5.active
ws5.title = "Danh Sách Phòng Thi"

ws5.merge_cells("A1:H1")
ws5["A1"] = "DANH SÁCH THÍ SINH VÀ PHIẾU ĐIỂM PHÒNG THI"
ws5["A1"].font = title_font
ws5["A1"].alignment = align_center

ws5.merge_cells("A2:H2")
ws5["A2"] = "Môn thi: IT301 - Lập trình Web   |   Học kỳ: Học kỳ 1 2026-2027"
ws5["A2"].font = subtitle_font
ws5["A2"].alignment = align_center

ws5.merge_cells("A3:H3")
ws5["A3"] = "Phòng thi: P.01 (Giảng đường A2 - Phòng 1)   |   Ngày thi: 15/12/2026   |   Giờ thi: 07:30 - 09:30"
ws5["A3"].font = subtitle_font
ws5["A3"].alignment = align_center

ws5.merge_cells("A4:H4")
ws5["A4"] = "Hướng dẫn: Nhập điểm thi thang 0-10 vào cột 'Điểm thi' (hoặc mã: V - Vắng thi, I - Hoãn thi). KHÔNG sửa các cột SBD, MSSV, Họ tên, Ngày sinh."
ws5["A4"].font = note_font
ws5["A4"].alignment = align_center

headers = [
    "STT", "Số báo danh (SBD)", "Mã sinh viên (MSSV)", "Họ và tên",
    "Ngày sinh", "Lớp sinh hoạt", "Điểm thi (0-10)", "Ghi chú / Kỷ luật"
]
ws5.append([]) # Empty row 5
ws5.append(headers) # Row 6

for col_idx in range(1, 9):
    c = ws5.cell(row=6, column=col_idx)
    c.fill = header_fill_blue
    c.font = header_font
    c.alignment = align_center
    c.border = cell_border
ws5.row_dimensions[6].height = 26

for idx, (sbd, code, name, dob, clazz, score, special, note) in enumerate(SAMPLE_STUDENTS, start=1):
    row_num = 6 + idx
    fill_to_use = zebra_fill if idx % 2 == 0 else white_fill
    
    score_val = special if special else score
    
    row_data = [idx, sbd, code, name, dob, clazz, score_val, note]
    ws5.append(row_data)
    
    for c_idx in range(1, 9):
        cell = ws5.cell(row=row_num, column=c_idx)
        cell.font = Font(name=FONT_NAME, size=11)
        cell.border = cell_border
        cell.fill = fill_to_use
        
        if c_idx in [1, 2, 3, 5, 6, 7]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            
        if c_idx == 7 and isinstance(score_val, float):
            cell.number_format = "0.0"
            
    ws5.row_dimensions[row_num].height = 20

# Column widths
widths = [8, 20, 16, 26, 14, 14, 16, 24]
for i, w in enumerate(widths, 1):
    ws5.column_dimensions[get_column_letter(i)].width = w

file5_path = os.path.join(OUTPUT_DIR, "05_PhongThi_HopLe_IT301_P01.xlsx")
wb5.save(file5_path)
print(f"Created: {file5_path}")

# =============================================================
# FILE 6: 06_PhongThi_CoLoi_DoiSoat.xlsx (Validation & Error Catching)
# Chứa 4 lỗi thực tế:
# 1. Dòng 5 (STT 5): Điểm vượt trần 11.5
# 2. Dòng 10 (STT 10): Điểm để trống
# 3. Dòng 15 (STT 15): Điểm nhập chữ lạ "Kém"
# 4. Dòng 20 (STT 20): Thí sinh sai MSSV "99999999" (không thuộc phòng)
# =============================================================
wb6 = openpyxl.Workbook()
ws6 = wb6.active
ws6.title = "Danh Sách Phòng Thi"

ws6.merge_cells("A1:H1")
ws6["A1"] = "DANH SÁCH THÍ SINH VÀ PHIẾU ĐIỂM PHÒNG THI"
ws6["A1"].font = title_font
ws6["A1"].alignment = align_center

ws6.merge_cells("A2:H2")
ws6["A2"] = "Môn thi: IT301 - Lập trình Web   |   Học kỳ: Học kỳ 1 2026-2027"
ws6["A2"].font = subtitle_font
ws6["A2"].alignment = align_center

ws6.merge_cells("A3:H3")
ws6["A3"] = "Phòng thi: P.01 (Giảng đường A2 - Phòng 1)   |   Ngày thi: 15/12/2026   |   Giờ thi: 07:30 - 09:30"
ws6["A3"].font = subtitle_font
ws6["A3"].alignment = align_center

ws6.merge_cells("A4:H4")
ws6["A4"] = "Hướng dẫn: Nhập điểm thi thang 0-10 vào cột 'Điểm thi' (hoặc mã: V - Vắng thi, I - Hoãn thi). KHÔNG sửa các cột SBD, MSSV, Họ tên, Ngày sinh."
ws6["A4"].font = note_font
ws6["A4"].alignment = align_center

ws6.append([]) # Row 5
ws6.append(headers) # Row 6

for col_idx in range(1, 9):
    c = ws6.cell(row=6, column=col_idx)
    c.fill = header_fill_blue
    c.font = header_font
    c.alignment = align_center
    c.border = cell_border
ws6.row_dimensions[6].height = 26

for idx, (sbd, code, name, dob, clazz, score, special, note) in enumerate(SAMPLE_STUDENTS, start=1):
    row_num = 6 + idx
    fill_to_use = zebra_fill if idx % 2 == 0 else white_fill
    
    curr_code = code
    curr_score = special if special else score
    is_err_row = False
    
    # Tạo 4 lỗi cố tình
    if idx == 5:
        curr_score = 11.5 # Lỗi điểm vượt trần
        is_err_row = True
    elif idx == 10:
        curr_score = None # Lỗi điểm để trống
        is_err_row = True
    elif idx == 15:
        curr_score = "Kém" # Lỗi ký tự lạ
        is_err_row = True
    elif idx == 20:
        curr_code = "99999999" # Lỗi không thuộc phòng thi
        is_err_row = True
        
    row_data = [idx, sbd, curr_code, name, dob, clazz, curr_score, note]
    ws6.append(row_data)
    
    for c_idx in range(1, 9):
        cell = ws6.cell(row=row_num, column=c_idx)
        cell.font = Font(name=FONT_NAME, size=11)
        cell.border = cell_border
        cell.fill = fill_to_use
        
        if c_idx in [1, 2, 3, 5, 6, 7]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            
        # Tô đỏ ô bị lỗi
        if idx == 5 and c_idx == 7:
            cell.fill = highlight_error_fill
            cell.font = highlight_error_font
        elif idx == 10 and c_idx == 7:
            cell.fill = highlight_error_fill
        elif idx == 15 and c_idx == 7:
            cell.fill = highlight_error_fill
            cell.font = highlight_error_font
        elif idx == 20 and c_idx == 3:
            cell.fill = highlight_error_fill
            cell.font = highlight_error_font
            
        if c_idx == 7 and isinstance(curr_score, float):
            cell.number_format = "0.0"
            
    ws6.row_dimensions[row_num].height = 20

for i, w in enumerate(widths, 1):
    ws6.column_dimensions[get_column_letter(i)].width = w

file6_path = os.path.join(OUTPUT_DIR, "06_PhongThi_CoLoi_DoiSoat.xlsx")
wb6.save(file6_path)
print(f"Created: {file6_path}")
